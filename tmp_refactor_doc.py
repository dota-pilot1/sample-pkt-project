from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

path = '/Users/terecal/pilot-project/sample-pkt-project/output/tikitaka-three-goals.docx'
doc = Document(path)
replacements = {
    '티키타카 노트의 3대 목표': '티키타카 노트의 세 가지 개발 목표',
    '사람과 Agent가 함께 읽고 활용하는 개발 지식': '사람이 배우고 Agent와 함께 만드는 개발 지식',
    '티키타카 노트는 개발 과정의 자료를 문서로 쌓고, Agent가 API로 읽고 편집할 수 있도록 연결하는 개발 노트 앱입니다. 학습과 Agent 활용을 함께 이어 가며, 축적한 지식을 사내 개발 정보 시스템으로 확장하는 것을 목표로 합니다.': '티키타카 노트는 개발하면서 배운 내용과 작업 기록을 한곳에 남기는 앱입니다. 사람은 원리와 설계를 이해하고, Agent는 문서를 읽고 정리하고 수정하는 일을 돕습니다. 이렇게 쌓인 지식을 다음 개발에 다시 쓰고, 나중에는 사내 개발 지식 시스템으로 넓히는 것이 이 앱을 만든 이유입니다.',
    '01 Agent를 활용한 문서화와 각종 자료의 API 제공': '01 Agent와 함께하는 문서 작성',
    '문서·작성 예제·구현 기록을 Agent가 조회하고 활용하도록 연결합니다. 필요한 API와 자료를 골라 전달하는 작업 도우미가 출발점입니다.': '개발 문서와 작성 예제를 Agent가 읽고 활용할 수 있게 연결합니다. 필요한 자료를 골라 Agent에게 전달하는 작업 도우미가 출발점입니다.',
    '02 학습 50과 Agent 활용 50의 개발 추구': '02 기본기 학습과 Agent 활용의 균형',
    '사람이 원리를 이해하고 판단하며, Agent와 함께 구현하고 정리합니다. 노트에 남긴 설계와 검증 기록으로 자신의 이해를 확인합니다.': '사람은 기술 원리와 설계를 이해하고 판단합니다. Agent는 구현과 정리, 검증을 돕습니다. 노트에는 이해한 내용과 실제로 확인한 결과를 함께 남깁니다.',
    '03 사내 개발 정보 시스템으로 확장': '03 사내 개발 지식 시스템으로 확장',
    '프로젝트 규칙과 설계, 구현 패턴, 실패 사례를 축적하고, 향후 RAG를 통해 Agent가 사내 지식을 근거로 활용하도록 확장합니다.': '프로젝트 규칙과 설계 이유, 해결 방법, 실패 사례를 모읍니다. 앞으로는 필요한 자료를 검색해 사람과 Agent가 함께 참고하는 사내 개발 지식 시스템으로 확장합니다.',
    'Agent를 위한 문서화와 자료 API': 'Agent와 함께하는 문서 작성',
    '문서를 사람이 읽는 화면으로 제공하면서, 같은 내용을 Agent가 조회하고 수정할 수 있는 API로 연결합니다. 개발자는 작업에 필요한 자료와 요청 방식을 골라 Agent에게 전달합니다.': '문서는 사람이 읽기 쉬워야 하고, Agent도 읽고 활용할 수 있어야 합니다. 그래서 문서와 작성 예제를 API로 연결하고, 필요한 자료를 골라 Agent에게 전달할 수 있게 합니다.',
    '학습 50과 Agent 활용 50': '기본기 학습과 Agent 활용의 균형',
    '사람이 기본기를 공부하고 설계를 판단하는 과정과, Agent를 활용해 구현하고 검증을 돕는 과정을 균형 있게 가져갑니다. 티키타카 노트는 그 과정에서 이해한 내용과 실행한 결과를 함께 남기는 도구입니다.': 'Agent에게 모든 일을 맡기는 것이 목표는 아닙니다. 사람은 기술 원리와 설계를 이해하고, Agent는 예제 작성과 구현, 문서 정리를 돕습니다. 티키타카 노트에는 이해한 내용과 실제로 확인한 결과를 함께 남깁니다.',
    '티키타카 노트는 개발하면서 배운 내용과 작업 기록을 한곳에 남기는 앱입니다. 사람은 원리와 설계를 이해하고, Agent는 문서를 읽고 정리하고 수정하는 일을 돕습니다. 노트는 도메인 설계와 프로토타이핑에도 활용합니다. 이렇게 쌓인 지식을 다음 개발에 다시 쓰고, 나중에는 사내 개발 지식 시스템으로 넓히는 것이 이 앱을 만든 이유입니다.': '티키타카 노트는 개발하면서 배운 내용과 작업 기록을 한곳에 남기는 앱입니다. 사람은 원리와 설계를 이해하고, Agent는 문서를 읽고 정리하고 수정하는 일을 돕습니다. 노트는 도메인 설계와 프로토타이핑에도 활용합니다. 이렇게 쌓인 지식을 다음 개발에 다시 쓰고, 나중에는 사내 개발 지식 시스템으로 넓히는 것이 이 앱을 만든 이유입니다.',
    'Agent에게 모든 일을 맡기는 것이 목표는 아닙니다. 사람은 기술 원리와 설계를 이해하고, Agent는 예제 작성과 구현, 문서 정리를 돕습니다. 티키타카 노트에는 이해한 내용과 실제로 확인한 결과를 함께 남깁니다.': 'Agent에게 모든 일을 맡기는 것이 목표는 아닙니다. 사람은 기술 원리와 설계를 이해하고, Agent는 예제 작성과 구현, 문서 정리를 돕습니다. 티키타카 노트에는 이해한 내용과 실제로 확인한 결과를 함께 남깁니다. 또한 노트를 도메인 설계와 프로토타이핑에 활용해, 문제와 해결 방법을 개발 초기에 정리합니다.',
    '사내 개발 정보 시스템으로 확장': '사내 개발 지식 시스템으로 확장',
    '팀의 개발 규칙과 설계 결정, 구현 패턴, 실패 사례를 축적하고, 사람과 Agent가 필요한 때 찾아 활용하는 사내 개발 정보 시스템으로 확장할 예정입니다.': '개발하면서 정리한 규칙, 설계 이유, 해결 방법, 실패 사례를 계속 모읍니다. 앞으로는 이 자료를 검색해 사람과 Agent가 필요한 순간에 참고할 수 있는 사내 개발 지식 시스템으로 확장할 계획입니다.',
    '개발하면서 정리한 규칙, 설계 이유, 해결 방법, 실패 사례를 계속 모읍니다. 앞으로는 이 자료를 검색해 사람과 Agent가 필요한 순간에 참고할 수 있는 사내 개발 지식 시스템으로 확장할 계획입니다.': '개발하면서 정리한 규칙과 해결 방법을 모읍니다. 앞으로는 이 자료를 검색해 사람과 Agent가 필요한 순간에 참고할 수 있는 사내 개발 지식 시스템으로 확장할 계획입니다.',
    '전달한 API 안내로 현재 주제의 본문과 하위 문서를 먼저 읽어줘. 선택한 작성 예제를 참고해 프로토타입의 목적, 설계 이유, 라이브러리 사용법, 검증 결과를 정리해줘. 기존 내용과 문서 구조를 확인한 뒤 필요한 부분만 추가하거나 수정하고, 실행하지 않은 검증은 미실행으로 표시해줘. 저장 후에는 문서 위치와 본문을 다시 확인해줘. 등등': '전달한 API 안내로 현재 주제의 본문과 하위 문서를 먼저 읽어줘. 선택한 작성 예제를 참고해 프로토타입의 목적, 설계 이유, 라이브러리 사용법, 검증 결과를 정리해줘. 기존 내용과 문서 구조를 확인한 뒤 필요한 부분만 추가하거나 수정하고, 실행하지 않은 검증은 미실행으로 표시해줘. 저장 후에는 문서 위치와 본문을 다시 확인해줘. 필요한 내용이 있으면 마지막에 덧붙여줘.',
}
for para in doc.paragraphs:
    if para.text in replacements:
        para.text = replacements[para.text]

# Make the opening summary match the requested four-part purpose statement.
for para in doc.paragraphs:
    if para.text == '티키타카 노트의 세 가지 개발 목표':
        para.text = '티키타카 노트 앱을 만든 이유'
    elif para.text == '사람이 배우고 Agent와 함께 만드는 개발 지식':
        para.text = '노트를 쓰고, 배우고, 설계하고, Agent와 함께 다시 활용하기'
    elif para.text == '티키타카 노트는 개발하면서 배운 내용과 작업 기록을 한곳에 남기는 앱입니다. 사람은 원리와 설계를 이해하고, Agent는 문서를 읽고 정리하고 수정하는 일을 돕습니다. 이렇게 쌓인 지식을 다음 개발에 다시 쓰고, 나중에는 사내 개발 지식 시스템으로 넓히는 것이 이 앱을 만든 이유입니다.':
        para.text = '티키타카 노트는 개발하면서 배운 내용과 작업 기록을 한곳에 남기는 앱입니다. 노트를 직접 편집하고, 개발 기본기를 배우고, 도메인을 설계하고, 프로토타입을 만들어 봅니다. Agent는 문서를 읽고 정리하고 수정하는 일을 돕습니다. 앞으로는 MCP와 RAG 등을 활용해 개발 지식 정보 시스템으로 확장하려고 합니다.'
    elif para.text == '티키타카 노트는 개발하면서 배운 내용과 작업 기록을 한곳에 남기는 앱입니다. 사람은 원리와 설계를 이해하고, Agent는 문서를 읽고 정리하고 수정하는 일을 돕습니다. 노트는 도메인 설계와 프로토타이핑에도 활용합니다. 이렇게 쌓인 지식을 다음 개발에 다시 쓰고, 나중에는 사내 개발 지식 시스템으로 넓히는 것이 이 앱을 만든 이유입니다.':
        para.text = '티키타카 노트는 개발하면서 배운 내용과 작업 기록을 한곳에 남기는 앱입니다. 노트를 직접 편집하고, 개발 기본기를 배우고, 도메인을 설계하고, 프로토타입을 만들어 봅니다. Agent는 문서를 읽고 정리하고 수정하는 일을 돕습니다. 앞으로는 MCP와 RAG 등을 활용해 개발 지식 정보 시스템으로 확장하려고 합니다.'

for para in doc.paragraphs:
    if para.text.startswith('티키타카 노트는 개발하면서 배운 내용과 작업 기록을 한곳에 남기는 앱입니다.'):
        para.text = '티키타카 노트는 개발하면서 배운 내용과 작업 기록을 한곳에 남기는 앱입니다. 노트를 직접 편집하고, 개발 기본기를 배우고, 도메인을 설계하고, 프로토타입을 만들어 봅니다. 프롬프트와 디자인 자료를 모아 다시 쓰고, 필요한 노트를 쉽게 검색할 수 있게 합니다. 이후에는 MCP와 RAG를 활용한 개발 지식 정보 시스템으로 확장하려고 합니다.'

def insert_before(target, text, style='Normal'):
    new = OxmlElement('w:p')
    target._element.addprevious(new)
    from docx.text.paragraph import Paragraph
    paragraph = Paragraph(new, target._parent)
    if isinstance(style, str):
        source = next((p for p in doc.paragraphs if p.style.name == style), target)
        paragraph.style = source.style
    else:
        paragraph.style = style
    paragraph.add_run(text)
    return paragraph

# Add domain design and prototyping as its own purpose in the detailed section.
paras = doc.paragraphs
if not any(p.text == '도메인 설계와 프로토타이핑' for p in paras):
    target = next((p for p in paras if p.text == '목표 03  향후 확장 계획'), None)
    if target:
        insert_before(target, '문제와 데이터를 먼저 정리하고, 작은 화면과 기능을 만들어 보며 설계를 확인합니다. 노트에 남긴 초안과 결과를 바탕으로 다음 개발 방향을 정합니다.', 'Normal')
        insert_before(target, '목표 03', 'Normal')
        insert_before(target, '도메인 설계와 프로토타이핑', 'Heading 1')
        target.text = '목표 04  향후 확장 계획'

# Add the domain/prototyping item to the opening summary.
if not any(p.text == '03 도메인 설계와 프로토타이핑에 활용' for p in doc.paragraphs):
    target = next((p for p in doc.paragraphs if p.text == '03 사내 개발 지식 시스템으로 확장'), None)
    if target:
        insert_before(target, '도메인과 프로토타입을 먼저 정리하고, 작은 화면과 기능을 만들어 보며 설계를 확인합니다.', 'Normal')
        insert_before(target, '03 도메인 설계와 프로토타이핑에 활용', 'Heading 2')
        target.text = '04 프롬프트 라이브러리와 디자인 시스템으로 활용'
        for p in doc.paragraphs:
            if p.text == '프로젝트 규칙과 설계 이유, 해결 방법, 실패 사례를 모읍니다. 앞으로는 필요한 자료를 검색해 사람과 Agent가 함께 참고하는 사내 개발 지식 시스템으로 확장합니다.':
                p.text = '잘 만든 프롬프트와 공통 컴포넌트, UI/UX 참고 자료를 모아 다시 활용합니다.'

if not any(p.text == '05 모든 노트를 쉽게 검색하는 기능' for p in doc.paragraphs):
    target = next((p for p in doc.paragraphs if p.text == '04 프롬프트 라이브러리와 디자인 시스템으로 활용'), None)
    if target:
        insert_before(target, '쌓인 노트를 빠르게 찾고, 필요한 개발 지식과 작업 기록을 다시 활용합니다.', 'Normal')
        insert_before(target, '05 모든 노트를 쉽게 검색하는 기능', 'Heading 2')

# Normalize the four summary paragraphs into heading and explanation pairs.
summary = doc.paragraphs
needed = {
    '04 프롬프트 라이브러리와 디자인 시스템으로 활용': None,
    '잘 만든 프롬프트와 공통 컴포넌트, UI/UX 참고 자료를 모아 다시 활용합니다.': None,
    '05 모든 노트를 쉽게 검색하는 기능': None,
    '쌓인 노트를 빠르게 찾고, 필요한 개발 지식과 작업 기록을 다시 활용합니다.': None,
}
for p in summary:
    if p.text in needed:
        needed[p.text] = p._element
if all(needed.values()):
    body = needed['04 프롬프트 라이브러리와 디자인 시스템으로 활용'].getparent()
    anchor = needed['04 프롬프트 라이브러리와 디자인 시스템으로 활용']
    pos = body.index(anchor)
    for element in needed.values():
        body.remove(element)
    order = [
        needed['04 프롬프트 라이브러리와 디자인 시스템으로 활용'],
        needed['잘 만든 프롬프트와 공통 컴포넌트, UI/UX 참고 자료를 모아 다시 활용합니다.'],
        needed['05 모든 노트를 쉽게 검색하는 기능'],
        needed['쌓인 노트를 빠르게 찾고, 필요한 개발 지식과 작업 기록을 다시 활용합니다.'],
    ]
    for offset, element in enumerate(order):
        body.insert(pos + offset, element)

for p in doc.paragraphs:
    if p.text == '04 개발 지식 정보 시스템으로 확장':
        p.text = '04 프롬프트 라이브러리와 디자인 시스템으로 활용'
    elif p.text == 'MCP와 RAG 등을 활용해 개발 지식을 검색하고, 사람과 Agent가 함께 참고하는 시스템으로 확장합니다.':
        p.text = '잘 만든 프롬프트와 공통 컴포넌트, UI/UX 참고 자료를 모아 다시 활용합니다.'
    elif p.text == '목표 04  향후 확장 계획':
        p.text = '이후 계획'

if not any(p.text == '05 모든 노트를 쉽게 검색하는 기능' for p in doc.paragraphs):
    target = next((p for p in doc.paragraphs if p.text == '04 프롬프트 라이브러리와 디자인 시스템으로 활용'), None)
    if target:
        insert_before(target, '쌓인 노트를 빠르게 찾고, 필요한 개발 지식과 작업 기록을 다시 활용합니다.', 'Normal')
        insert_before(target, '05 모든 노트를 쉽게 검색하는 기능', 'Heading 2')

for p in doc.paragraphs:
    if p.text == '목표 04  향후 확장 계획':
        p.text = '이후 계획'

# Keep the summary heading before its explanation after XML insertion.
summary = doc.paragraphs
for i in range(len(summary) - 1):
    if summary[i].text == '도메인과 프로토타입을 먼저 정리하고, 작은 화면과 기능을 만들어 보며 설계를 확인합니다.' and summary[i + 1].text == '03 도메인 설계와 프로토타이핑에 활용':
        old_style = summary[i].style
        summary[i].text, summary[i + 1].text = summary[i + 1].text, summary[i].text
        summary[i].style, summary[i + 1].style = summary[i + 1].style, old_style
        break

# Remove trailing manual line breaks left by the source document so they do not
# create an otherwise empty final page after the wording changes.
for para in doc.paragraphs:
    if para.text.startswith('\n사내 지식의 검색 범위'):
        para.text = para.text.strip()
    if '등등' in para.text:
        para.text = para.text.replace('등등', '필요한 내용이 있으면 마지막에 덧붙여줘')
for para in list(doc.paragraphs)[::-1]:
    if not para.text.strip():
        para._element.getparent().remove(para._element)
    else:
        break

# LibreOffice headless rendering needs an explicitly embedded East Asian font
# hint; otherwise the existing document's Korean text can render as blanks.
font_name = 'Arial Unicode MS'
for style in doc.styles:
    if hasattr(style, 'font'):
        style.font.name = font_name
        style._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
for para in doc.paragraphs:
    for run in para.runs:
        run.font.name = font_name
        run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.name = font_name
                    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
doc.save(path)
