from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = '/Users/terecal/pilot-project/sample-pkt-project/배포 관련 문서/PKT Study Fullstack 소개 문서.docx'
BLUE = RGBColor(31, 78, 121); MUTED = RGBColor(92, 104, 116); INK = RGBColor(35, 42, 50)

def set_font(run, size=10.5, bold=False, color=INK):
    run.font.name='AppleMyungjo'
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'AppleMyungjo')
    run.font.size=Pt(size); run.bold=bold; run.font.color.rgb=color

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)

def cell_margins(cell):
    tcPr=cell._tc.get_or_add_tcPr(); mar=OxmlElement('w:tcMar')
    for name,val in [('top',90),('start',130),('bottom',90),('end',130)]:
        x=OxmlElement('w:'+name); x.set(qn('w:w'),str(val)); x.set(qn('w:type'),'dxa'); mar.append(x)
    tcPr.append(mar)

def cell_width(cell, width):
    tcPr=cell._tc.get_or_add_tcPr(); w=OxmlElement('w:tcW'); w.set(qn('w:w'),str(width)); w.set(qn('w:type'),'dxa'); tcPr.append(w)

def hyperlink(p, text, url):
    rid=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),rid); r=OxmlElement('w:r'); rp=OxmlElement('w:rPr')
    c=OxmlElement('w:color'); c.set(qn('w:val'),'0563C1'); rp.append(c)
    u=OxmlElement('w:u'); u.set(qn('w:val'),'single'); rp.append(u); r.append(rp)
    t=OxmlElement('w:t'); t.text=text; r.append(t); h.append(r); p._p.append(h)

def para(doc, text='', style=None, size=10.5, bold=False, color=INK, before=0, after=5, align=None):
    p=doc.add_paragraph(style=style); p.paragraph_format.space_before=Pt(before); p.paragraph_format.space_after=Pt(after); p.paragraph_format.line_spacing=1.1
    if align is not None: p.alignment=align
    if text: set_font(p.add_run(text),size,bold,color)
    return p

doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
sec.top_margin=Inches(.68); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.78); sec.right_margin=Inches(.78)
sec.header_distance=Inches(.3); sec.footer_distance=Inches(.3)
normal=doc.styles['Normal']; normal.font.name='AppleMyungjo'; normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'AppleMyungjo'); normal.font.size=Pt(10.5); normal.font.color.rgb=INK
for name,size,color,before,after in [('Heading 1',16,BLUE,12,5),('Heading 2',12.5,BLUE,8,4)]:
    s=doc.styles[name]; s.font.name='AppleMyungjo'; s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'AppleMyungjo'); s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=color; s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after); s.paragraph_format.keep_with_next=True
hp=sec.header.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT; set_font(hp.add_run('PKT Study Fullstack  |  소개 문서'),9,False,MUTED)
fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT; set_font(fp.add_run('페이지 '),9,False,MUTED); fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); fp._p.append(fld)

para(doc,'PKT Study Fullstack',size=24,bold=True,color=BLUE,after=2)
para(doc,'PKT 학습 노트를 위한 독립형 데스크톱 애플리케이션',size=12,color=MUTED,after=10)
para(doc,'한눈에 보기',style='Heading 1')
para(doc,'PKT Study Fullstack은 Next.js, SQLite, TanStack Query, Lexical, Tauri를 결합한 설치형 학습 노트 앱입니다. 웹 브라우저에서 사용하는 서비스가 아니라, 학습 화면과 로컬 노트 데이터를 함께 제공하는 데스크톱 앱을 목표로 합니다.',after=6)

table=doc.add_table(rows=0,cols=2); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
for k,v in [('앱 이름','PKT Study Fullstack (PKT Study)'),('주요 용도','PKT 학습 노트 열람·작성 및 UI/개발 예제 학습'),('배포 방식','GitHub Releases를 통한 macOS·Windows 설치 파일 배포'),('데이터 저장','사용자 PC의 SQLite 데이터베이스'),('운영 API','https://api.hibot-docu.com/api')]:
    cs=table.add_row().cells
    for i,c in enumerate(cs): cell_width(c,2200 if i==0 else 7160); cell_margins(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    shade(cs[0],'EAF1F8'); set_font(cs[0].paragraphs[0].add_run(k),10,True,BLUE); set_font(cs[1].paragraphs[0].add_run(v),10)

para(doc,'공식 주소',style='Heading 1')
for label,url in [('소스 코드 저장소','https://github.com/dota-pilot1/pkt-study-fullstack'),('최신 릴리즈 및 설치 파일','https://github.com/dota-pilot1/pkt-study-fullstack/releases/latest'),('전체 릴리즈 목록','https://github.com/dota-pilot1/pkt-study-fullstack/releases'),('이슈·개선 제안','https://github.com/dota-pilot1/pkt-study-fullstack/issues')]:
    p=para(doc,after=2); set_font(p.add_run(label+'  '),10,True); hyperlink(p,url,url)

para(doc,'주요 기능과 특징',style='Heading 1')
for x in ['Lexical 기반의 구조화된 학습 노트 편집 및 코드 블록 표현','학습 주제와 문서를 계층적으로 탐색하는 PKT 플레이북 구조','Tauri 기반 설치형 앱으로 네트워크가 불안정한 환경에서도 로컬 데이터 확인 가능','GitHub Release에서 버전별 설치 파일과 변경 사항을 확인 가능']:
    p=para(doc,style='List Bullet',after=2); set_font(p.add_run(x))

para(doc,'데이터 보존 방식',style='Heading 1')
para(doc,'앱의 로컬 학습 데이터는 SQLite 파일로 관리됩니다. 개발 기준 DB는 프로젝트의 .data/pkt-study.db에 있으며, 릴리즈 빌드에서는 기준 시드로 패키징됩니다. 설치 앱은 사용자 데이터 경로에 DB를 복사해 사용하므로 앱 번들이 교체되어도 사용자 데이터가 분리되어 보존됩니다.',after=5)
para(doc,'다만 현재 배포 정책은 온보딩·학습용 테스트 앱을 전제로 합니다. 릴리즈 버전이 바뀔 때 기준 학습 데이터와 사용자 DB를 동기화할 수 있으므로, 중요한 개인 노트는 앱 내부 SQLite 백업 기능으로 먼저 백업하는 것이 좋습니다.',after=7)

para(doc,'설치 및 업데이트 안내',style='Heading 1')
for x in ['위의 “최신 릴리즈 및 설치 파일” 주소를 엽니다.','사용 중인 운영체제에 맞는 설치 파일을 내려받습니다.','설치 후 앱을 실행하고, 필요하면 앱 내부 설정의 SQLite 백업·복원 기능을 사용합니다.','새 버전이 나오면 같은 최신 릴리즈 주소에서 변경 사항과 새 설치 파일을 확인합니다.']:
    p=para(doc,style='List Number',after=2); set_font(p.add_run(x))

para(doc,'배포·기술 참고 문서',style='Heading 1')
para(doc,'릴리즈 절차, 버전 동기화, Windows Actions, macOS 빌드·서명, SQLite 기준 시드 검증은 프로젝트의 배포 문서에 정리되어 있습니다.',after=3)
p=para(doc,after=2); set_font(p.add_run('관련 문서 위치: '),10,True); set_font(p.add_run('배포 관련 문서/08-PKT-Study-Fullstack-하이브리드-릴리즈.md'),10,False,MUTED)
para(doc,'문의·참여',style='Heading 1')
para(doc,'사용 중 발견한 오류나 개선 아이디어는 GitHub Issues에 등록할 수 있습니다. 소스 코드와 릴리즈 기록은 공개 저장소에서 확인할 수 있으며, 각 버전의 설치 파일은 Releases 페이지에서 제공합니다.',after=0)
doc.save(OUT); print(OUT)
