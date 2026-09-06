from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path('/Users/terecal/pilot-project/sample-pkt-project')
OUT=ROOT/'output/documents/tikitaka-three-goals.docx'
SCREEN=Path('/var/folders/4h/z2zcntnn26v1szhfnc6m0t3w0000gn/T/codex-clipboard-09218608-3633-41b1-93fd-4d463a0d592d.png')
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.65)
sec.left_margin=sec.right_margin=Inches(.7)

def set_font(target,size=11,bold=False):
    target.font.name='Noto Sans KR';target.font.size=Pt(size);target.font.bold=bold
    target.font.color.rgb=RGBColor.from_string('000000')
    rf=target._element.get_or_add_rPr().get_or_add_rFonts()
    for key in list(rf.attrib):
        if key.endswith('Theme'):del rf.attrib[key]
    for key in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+key),'Noto Sans KR')

for name,size,bold in [('Normal',11,False),('Title',26,True),('Subtitle',14,False),('Heading 1',21,True),('Heading 2',13,True),('Heading 3',11,True)]:
    s=doc.styles[name];set_font(s,size,bold)
    s.paragraph_format.line_spacing=Pt(size*1.5)
    s.paragraph_format.space_before=Pt(0)
    s.paragraph_format.space_after=Pt(8)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(14)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(13)
doc.styles['Heading 2'].paragraph_format.space_after=Pt(7)
for e in doc.styles.element.xpath('.//w:pBdr'):e.getparent().remove(e)
doc.core_properties.title='티키타카 노트의 3대 목표'
doc.core_properties.subject='Agent를 위한 문서화와 API 및 학습과 사내 지식 활용'
doc.core_properties.author='티키타카 노트'
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run();set_font(r,9)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');r._r.addnext(f)

def p(text,size=11,bold=False,after=8):
    para=doc.add_paragraph();r=para.add_run(text);set_font(r,size,bold)
    para.paragraph_format.space_after=Pt(after)
    para.paragraph_format.line_spacing=Pt(size*1.5)
    return para

def page(label,title,first=False):
    if not first:doc.add_page_break()
    p(label,9.5,True,6)
    doc.add_paragraph(title,'Title' if first else 'Heading 1')

def h(text):doc.add_paragraph(text,'Heading 2')
def lead(text):p(text,12,True,12)

def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    for cells,vals in [(t.rows[0].cells,headers)]+[(t.add_row().cells,row) for row in rows]:
        for cell,txt,w in zip(cells,vals,widths):cell.text=txt;cell.width=Inches(w)
    borders=OxmlElement('w:tblBorders')
    for name in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+name)
        for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:e.set(qn('w:'+k),v)
        borders.append(e)
    t._tbl.tblPr.append(borders)
    for i,row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if i==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for j,cell in enumerate(row.cells):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'203B54' if i==0 else ('F2F5F8' if i%2==0 else 'FFFFFF'));pr.append(sh)
            margins=OxmlElement('w:tcMar')
            for name,value in [('top','90'),('bottom','90'),('left','110'),('right','110')]:
                e=OxmlElement('w:'+name);e.set(qn('w:w'),value);e.set(qn('w:type'),'dxa');margins.append(e)
            pr.append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(0);para.paragraph_format.line_spacing=Pt(14)
                for r in para.runs:
                    set_font(r,10.3,i==0)
                    r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '222222')
    p('',5,after=0)

def link(label,url):
    para=doc.add_paragraph();para.paragraph_format.space_after=Pt(4)
    x=OxmlElement('w:hyperlink');x.set(qn('r:id'),para.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=para.add_run(label);set_font(r,10);r.font.color.rgb=RGBColor.from_string('245A85');x.append(r._r);para._p.append(x)

page('앱 소개','티키타카 노트의 3대 목표',True)
doc.add_paragraph('사람과 Agent가 함께 읽고 활용하는 개발 지식','Subtitle')
p('2026년 9월 4일',9,after=18)
lead('티키타카 노트는 개발 과정의 자료를 문서로 쌓고, Agent가 API로 읽고 편집할 수 있도록 연결하는 개발 노트 앱입니다. 학습과 Agent 활용을 함께 이어 가며, 축적한 지식을 사내 개발 정보 시스템으로 확장하는 것을 목표로 합니다.')
h('01 Agent를 활용한 문서화와 각종 자료의 API 제공')
p('문서·작성 예제·구현 기록을 Agent가 조회하고 활용하도록 연결합니다. 필요한 API와 자료를 골라 전달하는 작업 도우미가 출발점입니다.')
h('02 학습 50과 Agent 활용 50의 개발 추구')
p('사람이 원리를 이해하고 판단하며, Agent와 함께 구현하고 정리합니다. 노트에 남긴 설계와 검증 기록으로 자신의 이해를 확인합니다.')
h('03 사내 개발 정보 시스템으로 확장')
p('프로젝트 규칙과 설계, 구현 패턴, 실패 사례를 축적하고, 향후 RAG를 통해 Agent가 사내 지식을 근거로 활용하도록 확장합니다.')

page('목표 01  현재 기능을 기반으로 확대','Agent를 위한 문서화와 자료 API')
lead('문서를 사람이 읽는 화면으로 제공하면서, 같은 내용을 Agent가 조회하고 수정할 수 있는 API로 연결합니다. 개발자는 작업에 필요한 자료와 요청 방식을 골라 Agent에게 전달합니다.')
p('아래의 2차 주제 노트 작업 도우미는 이 목표를 보여 주는 화면입니다. 왼쪽에서 작업에 필요한 API를 고르고, 오른쪽에서 문서 작성에 참고할 예제를 선택합니다. 선택한 항목을 복사해 Agent에게 전달하면, Agent는 현재 문서 구조와 작성 형식을 읽고 요청받은 작업을 수행하는 데 활용할 수 있습니다.')
img=doc.add_paragraph();img.paragraph_format.space_after=Pt(8)
img.paragraph_format.line_spacing=1
inline=img.add_run().add_picture(str(SCREEN),width=Inches(7.1))
inline._inline.docPr.set('descr','2차 주제 노트 작업 도우미. 왼쪽 API 선택 목록, 오른쪽 작성 예제 목록과 본문, 상단 선택 항목 복사 버튼.')
p('앱 화면 예시. API 목록과 작성 예제를 함께 선택해 Agent에게 필요한 작업 맥락을 전달합니다.',9,after=7)
p('화면에 표시된 주제 번호와 주소는 예시이며, 실제 작업에서는 현재 선택한 위치의 안내를 사용합니다.',9)

page('목표 01  활용 방법','API와 작성 예제를 Agent에게 전달')
h('화면에서 하는 세 가지 선택')
for n,title,body in [
 ('1','왼쪽에서 작업에 맞는 API 선택','기본, 문서 작성, 구현 기록, 코드 리뷰 등의 탭으로 필요한 API를 고릅니다. 문서 구조 조회, 본문 조회, 새 문서 생성, 내용 수정, 하위 문서 추가처럼 작업 단위가 구분되어 있습니다.'),
 ('2','오른쪽에서 작성 예제 선택','API 구현 노트, 프론트 구현 노트, 개요 문서 등 참고할 예제를 고르고 본문을 확인합니다. Agent에게 문서의 형식과 구체적인 작성 기준을 전달하는 자료가 됩니다.'),
 ('3','선택 항목을 복사해 Agent에게 전달','선택 API 복사, 선택 샘플 복사, 선택 항목 복사 버튼을 목적에 맞게 사용합니다. 화면의 안내처럼 공통 규칙과 현재 앱의 작성 규칙, 선택한 API 및 예제 조회 주소를 작업 요청과 함께 전달합니다.')]:
    para=doc.add_paragraph();r=para.add_run(f'{n}  {title}  ');set_font(r,11,True);r=para.add_run(body);set_font(r)
h('자료를 API로 연결하는 이유')
table(['연결할 자료','Agent가 활용하는 방식'],[
 ['문서 구조와 현재 본문','작업 위치와 기존 내용을 확인하고, 중복 작성이나 문맥 누락을 줄입니다.'],
 ['작성 예제와 구현 기록','팀이 사용하는 문서 형식과 코드 설명 방식, 검증 기록의 수준을 참고합니다.'],
 ['문서 생성과 편집 API','요청받은 내용을 노트에 반영하고, 저장한 결과를 다시 조회해 확인합니다.'],
 ['향후 연결할 개발 자료','설계 규칙, 라이브러리 적용 사례, 디버깅 기록 등으로 활용 범위를 넓힙니다.'],
],[1.65,5.45])
p('현재는 문서와 작성 예제 API를 바탕으로 이 흐름을 제공합니다. 자료 API를 확대하면 같은 자료를 사람이 앱에서 검토하고 Agent가 작업 중 참조할 수 있습니다. 외부 Agent가 앱의 API에 접근할 수 있는 환경에서 사용합니다.')
h('Agent에게 요청하는 예시')
p('전달한 API 안내로 현재 주제의 본문과 하위 문서를 먼저 읽어줘. 선택한 작성 예제를 참고해 프로토타입의 목적, 설계 이유, 라이브러리 사용법, 검증 결과를 정리해줘. 기존 내용과 문서 구조를 확인한 뒤 필요한 부분만 추가하거나 수정하고, 실행하지 않은 검증은 미실행으로 표시해줘. 저장 후에는 문서 위치와 본문을 다시 확인해줘.')
p('이 예시에는 작업 도우미에서 복사한 API와 작성 예제 안내를 함께 전달합니다. 수정 전 최신 문서 버전을 확인하고, 저장 후에는 앱에서 결과를 검토합니다.',9)

page('목표 02','학습 50과 Agent 활용 50')
lead('사람이 기본기를 공부하고 설계를 판단하는 과정과, Agent를 활용해 구현하고 검증을 돕는 과정을 균형 있게 가져갑니다. 티키타카 노트는 그 과정에서 이해한 내용과 실행한 결과를 함께 남기는 도구입니다.')
h('학습과 사람의 판단 50')
p('기술 원리, 데이터 흐름, 라이브러리 동작을 이해합니다. 설계 대안을 비교하고, 코드와 결과를 직접 확인하며, 선택 이유를 자신의 말로 설명합니다.')
h('Agent와 도구 활용 50')
p('예제 작성, 작은 단위 구현, 문서 정리, 검증 항목 도출에 Agent를 활용합니다. 관련 노트를 전달하고 결과를 확인하며 다음 작업에 재사용합니다.')
p('“50 대 50”은 학습과 툴링에 같은 비중을 둔다는 개발 원칙입니다. 참고 가이드의 Agent Engineering 방향처럼, 사람이 이해하고 검증할 수 있는 범위로 작업을 나누어 반복합니다.')
h('노트에 남기는 세 가지 개발 활동')
table(['개발 활동','노트에 남길 내용','확인할 질문'],[
 ['프로토타이핑','문제, 화면 초안, 사용 흐름, 완료 조건','이 화면은 어떤 문제를 해결하고 어떤 동작을 보여 주는가?'],
 ['설계 검증과 공유','API 계약, 데이터 흐름, 대안, 실행 근거','왜 이 설계를 선택했고 어떤 방법으로 확인했는가?'],
 ['라이브러리 활용 학습','최소 예제, 적용 조건, 오류와 해결 과정','이 코드를 설명하고 다른 기능에도 적용할 수 있는가?'],
],[1.4,2.45,3.25])
h('작은 기능 하나를 완성하는 흐름')
p('목표 정의 → 원리 학습 → Agent와 구현 → 사람의 검증 → 기록과 재사용',11,True)
p('예를 들어 LOT 상태 필터를 만든다면 허용되는 상태와 조회 조건을 먼저 이해합니다. Agent와 API 구현과 화면 구현을 작은 단위로 진행한 뒤, 실제 응답과 화면의 동작이 일치하는지 확인합니다. 마지막으로 설계 이유와 테스트 결과, 실패했던 시도를 노트에 남깁니다.')
p('한 작업의 본문에는 목표와 결론을 적고, 긴 구현 과정은 API와 Front 등의 하위 문서로 나눕니다. 다음 개발에서 관련 노트를 다시 읽거나 Agent에게 전달하면 학습한 내용을 실제 작업에 이어 쓸 수 있습니다.')

page('목표 03  향후 확장 계획','사내 개발 정보 시스템으로 확장')
lead('팀의 개발 규칙과 설계 결정, 구현 패턴, 실패 사례를 축적하고, 사람과 Agent가 필요한 때 찾아 활용하는 사내 개발 정보 시스템으로 확장할 예정입니다.')
p('프로젝트별 자료가 쌓이면 신규 개발자는 업무와 시스템의 구조를 이해하는 데 활용하고, 기존 개발자는 과거의 결정과 해결 사례를 다시 찾을 수 있습니다. Agent 역시 팀에서 검증한 규칙과 사례를 작업의 근거로 참조하도록 연결하는 것이 목표입니다.')
h('RAG로 필요한 사내 지식 연결')
p('RAG는 질문이나 작업과 관련된 문서를 검색하고, 검색한 내용을 근거로 답변을 생성하는 방식입니다. 향후에는 사람이 직접 골라 주는 자료에 더해, 작업과 관련된 문서를 검색해 Agent에게 공급하는 흐름을 추가할 계획입니다.')
p('검증된 노트 축적 → 관련 자료 검색 → Agent에 근거 전달 → 출처 확인과 문서 갱신',11,True)
p('“우리 프로젝트의 오류 처리 규칙에 맞춰 API를 설계해줘”라는 요청에 대해, 팀의 API 규칙과 대표 구현, 과거 오류 사례를 찾아 함께 제공하는 방식입니다. 답변에서 참조 문서와 기준 버전을 확인할 수 있도록 연결하고, 검증된 새 결과는 다시 노트에 반영하는 흐름을 지향합니다.')
table(['구분','현재의 출발점','향후 확장 방향'],[
 ['자료 전달','작업 도우미에서 API와 예제를 선택해 전달','RAG로 작업에 관련된 사내 문서 검색'],
 ['지식 활용','선택 문서와 예제를 바탕으로 작성·편집','사내 규칙과 사례를 근거로 답변·개발 지원'],
 ['관리 범위','앱의 노트와 자료 API','프로젝트별 지식, 접근 권한, 출처와 최신성 관리'],
],[1.1,2.95,3.05])
p('사내 지식의 검색 범위와 접근 권한, 변경·삭제된 문서의 반영, 문서 버전과 검증 시점 관리는 확장 시 함께 설계할 항목입니다. RAG와 사내 정보 시스템 연계는 향후 추가할 기능입니다.')
link('앱 설치와 최신 릴리즈','https://github.com/dota-pilot1/pkt-study-fullstack/releases/latest')
link('소스 코드와 이슈','https://github.com/dota-pilot1/pkt-study-fullstack')
p('참고 자료: 티키타카 개발 노트 Agent Engineering Guide의 학습과 툴링 균형, 작은 단위 검증, 노트 재사용 원칙. 화면 자료: 2차 주제 노트 작업 도우미.',9)

OUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUT)
print(OUT)
