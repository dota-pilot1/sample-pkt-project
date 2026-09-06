from pathlib import Path
from html import escape
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path('/Users/terecal/pilot-project/sample-pkt-project')
OUT = ROOT / 'output/pdf'
TITLE = '티키타카 노트 앱 소개'
SUBTITLE = '프로토타이핑과 설계 검증을 위한 개발 문서화'
PAGES = [
  {'id':'overview', 'title':TITLE, 'label':'앱 소개', 'blocks':[
    ('subtitle', SUBTITLE),
    ('meta', '2026년 9월 4일  |  앱 기준 버전 0.1.80'),
    ('lead', '티키타카 노트는 개발자가 아이디어를 구체화하고, 설계를 검증하고, 구현에서 배운 내용을 다시 활용하도록 돕는 데스크톱 개발 노트 앱입니다. 글과 코드, 구조도, UI 예제를 한 문맥 안에 정리하고 Agent와 함께 문서를 작성하고 개선할 수 있습니다.'),
    ('p', '개발 과정에서 필요한 것은 완성된 코드뿐 아니라 선택의 이유와 확인한 결과입니다. 티키타카 노트는 기능 단위로 목표와 구현 기록을 연결해, 다음 사람이 같은 문제를 이해하고 다음 Agent 작업에서도 필요한 맥락을 꺼내 쓸 수 있도록 문서화를 돕습니다.'),
    ('h2', '다섯 가지 활용 목적'),
    ('table', ['활용 목적','문서에 남기는 내용','기대 효과'], [
      ['프로토타이핑','화면 초안과 사용 흐름\n가정과 완료 조건','구현 전 아이디어 구체화'],
      ['설계 검증과 공유','API 계약과 데이터 흐름\n대안과 검증 근거','설계 의도와 판단 기준 공유'],
      ['라이브러리 활용 학습','최소 예제와 적용 조건\n소스와 실패 사례','실제 기능에 적용할 역량 축적'],
      ['Agent 기반 노트 편집','선택 문서의 문맥\n본문과 하위 문서','작성과 갱신 작업 지원'],
      ['RAG 기반\n사내 지식 활용','사내 규칙과 검증된 사례\n출처와 최신성 정보','Agent의 업무 맥락 공급 예정'],
    ]),
    ('h2', '학습과 도구 활용을 함께'),
    ('p', '참고 가이드의 “학습 50 + 툴링 50”은 기술 이해와 Agent 활용을 균형 있게 쌓자는 원칙입니다. 사람이 목표와 판단 기준을 정하고, Agent가 정리와 구현을 돕고, 사람이 실행 결과를 검증합니다. 이때 남긴 기록이 다음 작업의 출발점이 됩니다. [1]'),
    ('p', '기본 문서는 공간, 카테고리, 주제, 본문과 하위 문서로 정리합니다. 앱은 Tauri 기반으로 macOS와 Windows에서 제공되며, 로컬 SQLite 저장과 백업·복원을 지원합니다. RAG 연동은 향후 추가할 기능입니다. [2]'),
  ]},
  {'id':'design', 'title':'프로토타이핑과 설계 검증', 'label':'활용 1과 2', 'blocks':[
    ('lead', '문서에서 요구사항과 화면을 함께 살펴보고, 구현 전에 확인해야 할 질문을 드러냅니다. 작은 프로토타입과 검증 기록을 쌓으면 설계 회의와 개발 인수인계에 같은 자료를 사용할 수 있습니다.'),
    ('h2','프로토타이핑을 위한 문서화'),
    ('p','새 기능의 본문에는 해결할 문제, 사용자의 행동, 화면의 주요 상태와 완료 조건을 적습니다. UI 초안이나 코드 예제를 함께 두고 정상 상태뿐 아니라 빈 결과, 로딩, 오류 상태도 설명합니다. 화면 모양과 동작 의도를 함께 검토할 수 있어 구현 범위를 합의하기 쉽습니다.'),
    ('p','앱은 HTML 미리보기, Tailwind TSX 미리보기, 등록된 UI 컴포넌트 미리보기를 제공합니다. 지원되는 범위 안에서 화면 조각을 확인하고 문서의 설명과 비교할 수 있습니다. 미리보기는 UI 아이디어를 점검하는 용도이며, 실제 API 연동과 전체 애플리케이션 동작은 별도로 실행해 검증합니다. [2]'),
    ('h2','설계 검증과 공유를 위한 문서화'),
    ('p','설계 문서에는 API의 입력과 응답, 데이터의 이동 경로, 예외 처리, 선택한 방법과 대안을 남깁니다. Mermaid 구조도와 표, 코드 블록을 활용하면 화면·서버·데이터의 관계를 구체적으로 설명할 수 있습니다. 리뷰에서는 “왜 이 방법을 골랐는가”와 “어떻게 맞는지 확인했는가”를 함께 다룹니다.'),
    ('p','팀원이 같은 문서를 보며 검토할 수 있도록 본문에는 목표와 설계 결론을, 하위 문서에는 구현 상세와 검증 기록을 배치합니다. 이 구조는 회의 자료, 코드 리뷰 설명, 온보딩 자료에 활용할 수 있습니다. 여기서 공유는 문서를 공통 검토 자료로 쓰는 활용 방식이며, 실시간 공동 편집이나 중앙 지식 서버를 의미하지 않습니다.'),
    ('h2','활용 예시인 MES LOT 상태 필터'),
    ('table',['문서 단위','남길 내용'],[
      ['본문','상태 필터의 목적, 선택 가능한 값, 화면 초안, 완료 조건'],
      ['API 구현 하위 문서','쿼리 계약, 조회 조건, 잘못된 값 처리, 실제 요청과 응답'],
      ['Front 구현 하위 문서','선택 UI, URL 상태, 새로고침 유지, 빈 결과와 오류 화면'],
      ['검증 기록','API 계약과 화면 결과의 일치 여부, 테스트 명령, 실패 사례'],
    ]),
    ('p','한 TODO를 하나의 본문 문서로 두고, 긴 작업은 API 구현과 Front 구현 하위 문서로 나눌 수 있습니다. 프론트엔드만 필요한 작업은 UI 구현과 검증에 집중합니다. 참고 가이드의 LOT 예시도 계약 정의, 작은 단위 구현, 실행 검증, 노트 작성을 연결합니다. [1, 7쪽]'),
  ]},
  {'id':'libraries','title':'라이브러리 활용 능력을 쌓는 문서화','label':'활용 3','blocks':[
    ('lead','라이브러리를 이해했다는 기준을 실제 기능에 적용하고 그 선택을 설명할 수 있는 상태로 둡니다. 설치 방법, 최소 예제, 프로젝트 적용 조건, 검증 결과를 함께 기록해 학습을 다음 구현에 연결합니다.'),
    ('h2','예제를 실행하고 소스와 연결'),
    ('p','UI 갤러리에서는 등록된 컴포넌트의 속성을 바꿔 결과를 확인하고 소스를 살펴볼 수 있습니다. 글로 적은 사용법과 실제 표현을 함께 비교하면서 어떤 입력이 어떤 화면을 만드는지 학습합니다. 같은 컴포넌트를 어디에 재사용할지, 어떤 상태를 추가로 확인해야 할지도 문서에 남깁니다. [2]'),
    ('table',['학습 주제 예시','노트에서 확인할 질문'],[
      ['React와 UI 컴포넌트','입력 속성과 상태는 무엇이며 화면은 어떻게 바뀌는가'],
      ['Tailwind CSS','간격과 색상, 반응형 규칙을 프로젝트에 어떻게 맞출 것인가'],
      ['TanStack Query','조회 키와 캐시, 갱신 시점, 로딩·오류 처리는 어떻게 설계할 것인가'],
      ['Lexical','본문과 코드 블록을 어떤 문서 구조로 저장하고 편집할 것인가'],
      ['Mermaid','흐름도에 어떤 단계와 의존 관계를 보여 줄 것인가'],
    ]),
    ('p','위 항목은 앱에서 활용하는 기술을 바탕으로 한 학습 주제 예시입니다. 모든 라이브러리를 노트 안에서 자유롭게 설치하거나 실행하는 기능을 뜻하지는 않습니다. 현재 Tailwind TSX 미리보기는 제한된 화면 샘플을 위한 것으로 외부 import나 네트워크 호출을 지원하지 않습니다. [2]'),
    ('h2','다음 작업에서도 쓸 수 있는 기록'),
    ('p','라이브러리 노트는 “적용하려는 문제와 사용 버전 → 최소 동작 예제 → 프로젝트 적용 방식 → 검증 결과 → 재사용 조건” 순서로 작성합니다. 예제 코드를 붙인 뒤에는 필요한 입력, 관련 파일, 예상 출력과 실제 결과를 설명합니다.'),
    ('p','잘되지 않은 시도도 자산이 됩니다. 재현 조건과 오류 메시지, 원인, 해결 과정, 선택하지 않은 대안을 남기면 다음 작업에서 같은 시행착오를 줄이는 데 도움이 됩니다. “언제 다시 쓰면 되는가”와 “어떤 조건에서는 바꿔야 하는가”까지 적는 방식은 참고 가이드의 노트 기본 틀과 연결됩니다. [1, 8쪽]'),
    ('h2','팀 학습에 활용'),
    ('p','신규 개발자는 대표 노트의 예제를 수정하고 결과를 설명하는 방식으로 학습할 수 있습니다. 리뷰어는 소스와 함께 적용 이유, 실패 경로, 검증 근거를 확인합니다. 이런 기록이 쌓이면 코드 패턴, UI 사용 규칙, 디버깅 사례를 프로젝트별 학습 자료로 정리할 수 있습니다.'),
  ]},
  {'id':'agent','title':'Agent를 활용한 노트 편집','label':'활용 4','blocks':[
    ('lead','앱의 API for LLM 안내를 Agent에 전달하면 선택한 문서의 맥락을 읽고 본문을 정리하거나 하위 문서를 추가하는 작업에 활용할 수 있습니다. 사용자는 편집 목적과 범위를 정하고, 변경된 문서와 실행 근거를 확인합니다.'),
    ('h2','문서의 위치와 문맥을 함께 전달'),
    ('p','앱은 문서 조회, 본문과 하위 문서의 문맥 조회, 문서 생성과 편집을 위한 API 안내를 제공합니다. 현재 앱 화면에 표시된 주소와 요청 예제를 사용하면 문서 위치와 연결 정보를 Agent에게 전달할 수 있습니다. 외부 Agent가 해당 API에 접근할 수 있는 실행 환경이 필요합니다. [2]'),
    ('steps',[
      ('대상 확인','앱에서 편집할 주제나 본문을 선택하고 API 안내를 전달합니다. Agent는 tree 또는 topic 조회로 선택 위치를 확인합니다.'),
      ('기존 내용 읽기','본문과 하위 문서 문맥을 먼저 조회해 제목, 부모 문서, 기존 내용과 최신 버전을 확인합니다.'),
      ('범위에 맞춰 편집','설명 보완, 구현 기록 정리, 검증 항목 추가 등 요청 범위에 맞게 작성합니다. 본문은 앱의 Lexical 형식을 따릅니다.'),
      ('저장 결과 확인','조회한 버전을 expectedVersion으로 전달해 저장하고, 충돌하면 최신 문서를 다시 읽습니다. 저장 후 내용과 위치를 재조회합니다.'),
      ('사람이 검토','앱에서 본문과 코드 블록의 표시를 확인하고, 기술 설명과 실행 결과가 맞는지 판단합니다.'),
    ]),
    ('h2','Agent에게 요청하는 예시'),
    ('p','“선택한 LOT 상태 필터 문서와 하위 문서를 먼저 읽어줘. 기존 목표와 API 계약을 유지하면서 구현 이유, 빈 결과 처리, 검증 방법을 보완해줘. 아직 실행하지 않은 검증은 미실행으로 적고, 저장 후 문서 위치와 본문을 다시 확인해줘.”'),
    ('p','이 방식은 처음 문서를 만드는 일뿐 아니라 구현 변경 뒤 설명 갱신, 긴 기록의 분리, 누락된 검증 항목 점검에도 사용할 수 있습니다. Agent가 작성한 설명과 실제 코드·화면이 맞는지 확인하는 역할은 사람이 맡습니다.'),
    ('h2','지식을 다시 Agent 작업에 활용'),
    ('p','관련 노트에는 프로젝트 규칙, 대표 구현 방식, 과거 결정, 실패 경험이 함께 들어 있습니다. 현재 작업에 필요한 문서만 골라 Agent에게 주면 팀의 맥락을 전달할 수 있습니다. 참고 가이드가 제안한 “필요한 자료만 선택하고, 오래된 정보에는 기준 버전과 검증 시점을 남긴다”는 원칙을 적용합니다. [1, 8~9쪽]'),
  ]},
  {'id':'rag','title':'RAG 기반 사내 지식 활용 계획','label':'활용 5와 시작 안내','blocks':[
    ('lead','향후 티키타카 노트에 축적한 문서를 RAG와 연결해 Agent가 사내 규칙과 과거 사례를 참고하는 정보 시스템으로 확장할 예정입니다. RAG는 질문이나 작업과 관련된 문서를 검색한 뒤, 그 내용을 근거로 답변을 생성하는 방식입니다.'),
    ('p','예를 들어 “우리 프로젝트의 LOT 상태 변경 규칙에 맞춰 API를 설계해줘”라는 요청이 들어오면 관련 설계 노트, API 규칙, 과거 오류 사례를 찾아 Agent에게 제공하는 흐름을 목표로 합니다. 답변에는 참조 문서와 기준 버전을 연결해 사람이 근거를 확인할 수 있도록 할 계획입니다.'),
    ('h2','계획하는 정보 흐름'),
    ('p','문서와 메타데이터 정리 → 검색 인덱스 구성 → 작업에 필요한 근거 검색 → Agent 답변과 출처 연결 → 사람의 검증과 노트 갱신'),
    ('p','검색을 위해 문서를 작은 단위로 나누고 제목, 도메인, 기준 버전, 마지막 검증일을 연결하는 방안을 검토합니다. 사용자별 접근 범위와 변경·삭제 문서의 검색 반영, 근거가 부족한 경우의 응답 방식도 함께 설계합니다. 아래는 향후 구현 방향입니다.'),
    ('table',['구분','현재 제공','향후 확장 방향'],[
      ['문서 접근','선택한 문서와 문맥을 API로 조회','작업에 관련된 사내 문서 검색'],
      ['Agent 활용','전달받은 문맥으로 작성과 편집 지원','검색한 근거와 출처를 답변에 연결'],
      ['지식 관리','로컬 노트와 문서 구조 관리','최신성 관리와 사내 접근 권한 연계'],
    ]),
    ('h2','작은 기능 하나로 시작'),
    ('p','운영체제에 맞는 앱을 설치하고 샘플 노트를 읽은 뒤, 진행 중인 작은 기능 하나를 선택합니다. 본문에 목표와 완료 조건을 쓰고 UI 초안이나 최소 예제를 붙입니다. 구현과 검증 결과를 정리한 다음 Agent에게 설명 보완을 요청하고, 다음 작업에서 이 문서를 다시 참고합니다.'),
    ('p','활용 효과는 만든 문서의 수보다 기존 노트를 다시 사용한 사례, 별도 검증에서 통과한 작업, 변경자가 설계 이유를 설명할 수 있는지로 살펴볼 수 있습니다. 문서화가 실제 개발과 학습에 쓰이는지를 확인하는 기준입니다. [1, 11쪽]'),
    ('links',[
      ('앱 설치와 최신 릴리즈','https://github.com/dota-pilot1/pkt-study-fullstack/releases/latest'),
      ('소스 코드와 이슈','https://github.com/dota-pilot1/pkt-study-fullstack'),
    ]),
    ('source','참고 자료'),
    ('source','[1] 티키타카 개발 노트 Agent Engineering Guide, 총 12쪽. 학습과 도구 활용, 검증과 노트 재사용 원칙 참고. 파일: tikitaka-agent-engineering-guide.pdf.'),
    ('source','[2] pkt-study-fullstack 0.1.80 소스와 앱 API 안내. UI 갤러리, 미리보기, 문서 문맥 조회·편집 기능 및 릴리즈 구성의 기준 자료. 2026년 9월 4일 확인.'),
  ]},
]

def font_style(style, size, bold=False):
    style.font.name = 'Noto Sans KR'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0,0,0)
    rf = style.element.get_or_add_rPr().get_or_add_rFonts()
    for n in ['ascii','hAnsi','eastAsia','cs']:
        rf.set(qn('w:'+n), 'Noto Sans KR')
    for key in list(rf.attrib):
        if key.endswith('Theme'):del rf.attrib[key]
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.line_spacing = Pt(size * 1.5)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.65), Inches(.62)
sec.left_margin = sec.right_margin = Inches(.7)
for name,size,bold in [('Normal',11,False),('Title',27,True),('Subtitle',14,False),('Heading 1',21,True),('Heading 2',13,True)]:
    font_style(doc.styles[name],size,bold)
doc.styles['Title'].paragraph_format.space_after = Pt(9)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(15)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(13)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(6)
for element in doc.styles.element.xpath('.//w:pBdr'):
    element.getparent().remove(element)
doc.core_properties.title = TITLE
doc.core_properties.subject = SUBTITLE
doc.core_properties.author = '티키타카 노트'
doc.core_properties.keywords = '프로토타이핑, 설계 검증, 라이브러리, Agent, RAG'

footer=sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run=footer.add_run()
run.font.size=Pt(9)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');run._r.addnext(field)

def paragraph(text, size=None, bold=False, after=None):
    p=doc.add_paragraph()
    r=p.add_run(text);r.bold=bold
    if size:r.font.size=Pt(size)
    if after is not None:p.paragraph_format.space_after=Pt(after)
    return p

def table(headers, rows):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths=([1.47,3.22,2.41] if len(headers)==3 else [1.65,5.45])
    if headers==['구분','현재 제공','향후 확장 방향']:widths=[1.1,2.8,3.2]
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    for cell,w in zip(t.rows[0].cells,widths):cell.width=Inches(w)
    for c,h in zip(t.rows[0].cells,headers):c.text=h
    for row in rows:
        cells=t.add_row().cells
        for c,txt,w in zip(cells,row,widths):c.text=txt;c.width=Inches(w)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    t._tbl.tblPr.append(borders)
    for ri,row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if ri==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for ci,cell in enumerate(row.cells):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr()
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'203B54' if ri==0 else ('F2F5F8' if ri%2==0 else 'FFFFFF'));pr.append(sh)
            margins=OxmlElement('w:tcMar')
            for side,value in [('top','100'),('bottom','100'),('left','110'),('right','110')]:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),value);e.set(qn('w:type'),'dxa');margins.append(e)
            pr.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(14)
                if ci==0:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.size=Pt(10.3);r.bold=(ri==0)
                    r.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '222222')
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

def hyperlink(label,url):
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(4)
    h=OxmlElement('w:hyperlink')
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'245A85');pr.append(color);r.append(pr)
    txt=OxmlElement('w:t');txt.text=label;r.append(txt);h.append(r);p._p.append(h)

html_sections=[]
for page_no,page in enumerate(PAGES,1):
    if page_no>1:doc.add_page_break()
    paragraph(page['label'],9.5,bold=True,after=6)
    doc.add_paragraph(page['title'], 'Title' if page_no==1 else 'Heading 1')
    parts=[f'<section id="{page["id"]}" aria-labelledby="h-{page["id"]}"><div class="eyebrow">{escape(page["label"])}</div><h1 id="h-{page["id"]}">{escape(page["title"])}</h1>']
    for b in page['blocks']:
        kind=b[0]
        if kind=='table':
            table(b[1],b[2]);parts.append('<div class="table-scroll"><table><thead><tr>'+''.join('<th scope="col">'+escape(x)+'</th>' for x in b[1])+'</tr></thead><tbody>'+''.join('<tr>'+''.join('<td>'+escape(x).replace('\n','<br>')+'</td>' for x in row)+'</tr>' for row in b[2])+'</tbody></table></div>')
        elif kind=='steps':
            parts.append('<ol class="steps">')
            for i,(label,txt) in enumerate(b[1],1):
                p=doc.add_paragraph();p.paragraph_format.space_after=Pt(10)
                p.add_run(f'{i}  {label}  ').bold=True;p.add_run(txt)
                parts.append(f'<li><strong>{escape(label)}</strong><p>{escape(txt)}</p></li>')
            parts.append('</ol>')
        elif kind=='links':
            for label,url in b[1]:hyperlink(label,url);parts.append(f'<p class="link"><a href="{escape(url)}">{escape(label)} ↗</a></p>')
        elif kind=='h2':doc.add_paragraph(b[1],'Heading 2');parts.append('<h2>'+escape(b[1])+'</h2>')
        elif kind=='subtitle':doc.add_paragraph(b[1],'Subtitle');parts.append('<p class="subtitle">'+escape(b[1])+'</p>')
        else:
            size={'meta':9,'source':8.5,'lead':12}.get(kind)
            p=paragraph(b[1],size,kind=='lead',6 if kind=='source' else None)
            if kind=='source':p.paragraph_format.line_spacing=Pt(11.5)
            parts.append(f'<p class="{kind}">{escape(b[1])}</p>')
    parts.append('</section>');html_sections.append('\n'.join(parts))

OUT.mkdir(parents=True,exist_ok=True)
doc.save(OUT/'tikitaka-app-introduction.docx')
nav=''.join(f'<a href="#{p["id"]}"><span>{i:02d}</span>{escape(p["title"] if i>1 else "앱 한눈에 보기")}</a>' for i,p in enumerate(PAGES,1))
css='''
:root{color-scheme:light;--ink:#18252e;--muted:#576772;--accent:#245a85;--line:#d9d9d9}
*{box-sizing:border-box}html{scroll-behavior:smooth;scroll-padding-top:25px}body{margin:0;background:#edf1f4;color:var(--ink);font-family:"Apple SD Gothic Neo","Malgun Gothic","Noto Sans KR",sans-serif;font-size:17px;line-height:1.85;word-break:keep-all;overflow-wrap:anywhere}
a{color:var(--accent);text-underline-offset:4px}a:focus-visible,button:focus-visible{outline:3px solid #4488b5;outline-offset:4px}
.layout{max-width:1330px;margin:0 auto;display:grid;grid-template-columns:260px minmax(0,1fr);gap:40px;padding:48px 36px 72px}
aside{position:sticky;top:38px;align-self:start}.brand{font-size:21px;font-weight:800;letter-spacing:-.6px;margin-bottom:5px;color:#000}.edition{font-size:12px;letter-spacing:1.5px;color:var(--muted);margin-bottom:35px}
nav{display:flex;flex-direction:column;gap:16px}nav a{display:flex;gap:13px;color:#384a58;text-decoration:none;font-size:14px;line-height:1.65}nav a:hover{color:var(--accent)}nav span{font-size:12px;font-weight:700;color:#708b9e;padding-top:2px}
button{font:inherit;font-size:13px;cursor:pointer;background:white;border:1px solid #cbd4db;padding:8px 20px;border-radius:5px;margin-top:35px;color:var(--ink)}
main{min-width:0}section{background:#fff;padding:46px 52px 50px;margin-bottom:24px;box-shadow:0 8px 35px #1b364908;scroll-margin-top:25px}.eyebrow{font-size:12px;font-weight:700;letter-spacing:1px;color:#000;margin-bottom:13px}h1{font-size:30px;line-height:1.4;letter-spacing:-1.1px;color:#000;margin:0 0 22px}#overview h1{font-size:39px;margin-bottom:7px}h2{font-size:21px;line-height:1.5;letter-spacing:-.5px;color:#000;margin:31px 0 10px}p{margin:0 0 17px}.subtitle{font-size:20px;color:#000;margin-bottom:10px;line-height:1.6}.meta{font-size:12px;color:var(--muted);margin-bottom:30px}.lead{font-size:19px;font-weight:650;line-height:1.8;margin-bottom:23px}.source{font-size:12px;line-height:1.65;color:var(--muted);margin-bottom:7px}.link{font-size:14px;margin-bottom:4px}
.table-scroll{overflow-x:auto;margin:18px 0 24px}table{width:100%;border-collapse:collapse;font-size:14px;line-height:1.65}th,td{border:1px solid var(--line);padding:13px 14px;vertical-align:middle;text-align:left}th{background:#203b54;color:#fff;font-weight:650}tbody tr:nth-child(even){background:#f2f5f8}td:first-child{font-weight:650;min-width:102px}th:first-child{width:23%}.steps{padding-left:26px;margin:17px 0 27px}.steps li{padding-left:8px;margin-bottom:16px}.steps li::marker{color:var(--accent);font-weight:700}.steps p{margin:3px 0 0}
@media(max-width:950px){.layout{grid-template-columns:1fr;padding:24px 20px;gap:25px}aside{position:static}.edition{margin-bottom:16px}nav{flex-direction:row;flex-wrap:wrap;gap:12px 22px}nav a{font-size:13px}button{margin-top:20px}section{padding:36px 34px}}
@media(max-width:540px){body{font-size:16px}.layout{padding:20px 10px}aside{padding:0 14px}section{padding:28px 23px}#overview h1{font-size:31px}h1{font-size:27px}.lead{font-size:18px}.subtitle{font-size:18px}th,td{padding:10px;font-size:13px}table{min-width:460px}}
@media print{@page{size:letter;margin:17mm}body{background:white;font-size:10pt;line-height:1.5}.layout{display:block;padding:0;max-width:none}aside{display:none}section{box-shadow:none;padding:0;margin:0;break-after:page}section:last-child{break-after:auto}h1,#overview h1{font-size:23pt}h2{font-size:12pt;margin:16px 0 7px}.lead{font-size:11pt;line-height:1.55}.subtitle{font-size:13pt}.meta{margin-bottom:15px}p{margin-bottom:9px}table{font-size:9pt;min-width:0}th,td{padding:7px 9px;font-size:9pt}thead{display:table-header-group}tr{break-inside:avoid}.table-scroll{overflow:visible;margin:10px 0 15px}.steps li{margin-bottom:9px}.source{font-size:7.5pt}.link{font-size:9pt}a{color:#245a85}.eyebrow{font-size:8pt}}
'''
html='<!doctype html>\n<html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><meta name="description" content="티키타카 노트 앱 소개와 문서화 활용 사례 및 RAG 확장 계획"><title>'+TITLE+'</title><style>'+css+'</style></head><body><div class="layout"><aside><div class="brand">티키타카 노트</div><div class="edition">PRODUCT GUIDE · 2026.09</div><nav aria-label="문서 목차">'+nav+'</nav><button type="button" onclick="window.print()">인쇄 또는 PDF 저장</button></aside><main>'+''.join(html_sections)+'</main></div></body></html>'
(OUT/'tikitaka-app-introduction.html').write_text(html,encoding='utf-8')
print('Created',OUT/'tikitaka-app-introduction.docx')
print('Created',OUT/'tikitaka-app-introduction.html')
